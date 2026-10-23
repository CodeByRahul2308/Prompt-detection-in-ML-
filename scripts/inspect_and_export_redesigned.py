"""
Dump all slide content from Minor_Project_Presentation_Redesigned.pptx
and export slides to PNG images using PowerPoint COM.
"""
import os
import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

ppt_path = os.path.abspath("presentation/Minor_Project_Presentation_Redesigned.pptx")
out_img_dir = os.path.abspath("presentation/slide_images_redesigned")
os.makedirs(out_img_dir, exist_ok=True)

prs = Presentation(ppt_path)
print(f"Total slides: {len(prs.slides)}")
print(f"Dimensions: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} inches")

for idx, slide in enumerate(prs.slides, 1):
    print(f"\n{'='*60}\nSLIDE {idx}\n{'='*60}")
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                text = p.text.strip()
                if text:
                    font_sizes = [f"{run.font.size.pt:.1f}pt" for run in p.runs if run.font and run.font.size]
                    fs_str = f" [Fonts: {', '.join(font_sizes)}]" if font_sizes else ""
                    print(f"  P: {text}{fs_str}")
        elif shape.has_table:
            table = shape.table
            print(f"  TABLE ({len(table.rows)}x{len(table.columns)}):")
            for r_idx, row in enumerate(table.rows):
                row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                print(f"    R{r_idx}: {' | '.join(row_vals)}")
        else:
            print(f"  SHAPE: {shape.name} (type: {shape.shape_type})")

# Now export slides to PNG using PowerPoint COM
try:
    import win32com.client
    ppt_app = win32com.client.Dispatch("PowerPoint.Application")
    presentation = ppt_app.Presentations.Open(ppt_path, ReadOnly=True, Untitled=False, WithWindow=False)
    for i in range(1, presentation.Slides.Count + 1):
        img_file = os.path.join(out_img_dir, f"slide_{i:02d}.png")
        presentation.Slides(i).Export(img_file, "PNG", 1920, 1080)
        print(f"Exported slide {i} to {img_file}")
    presentation.Close()
    ppt_app.Quit()
    print("All slides exported successfully via win32com!")
except Exception as e:
    print(f"PowerPoint COM export via python failed ({e}), trying PowerShell...")
