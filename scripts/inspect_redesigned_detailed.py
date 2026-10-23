"""
Dump all slide content from Minor_Project_Presentation_Redesigned.pptx
to a UTF-8 text file and inspect slide details.
"""
import os
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

ppt_path = os.path.abspath("presentation/Minor_Project_Presentation_Redesigned.pptx")
prs = Presentation(ppt_path)

out_txt_path = os.path.abspath("presentation/slide_inspection_redesigned.txt")

with open(out_txt_path, "w", encoding="utf-8") as out:
    out.write(f"Total slides: {len(prs.slides)}\n")
    out.write(f"Dimensions: {prs.slide_width / 914400:.2f} x {prs.slide_height / 914400:.2f} inches\n")

    for idx, slide in enumerate(prs.slides, 1):
        out.write(f"\n{'='*70}\nSLIDE {idx}\n{'='*70}\n")
        for s_idx, shape in enumerate(slide.shapes):
            left_in = shape.left / 914400 if shape.left is not None else 0
            top_in = shape.top / 914400 if shape.top is not None else 0
            width_in = shape.width / 914400 if shape.width is not None else 0
            height_in = shape.height / 914400 if shape.height is not None else 0
            
            if shape.has_text_frame:
                out.write(f"\n  [Shape {s_idx}: {shape.name} at ({left_in:.2f}, {top_in:.2f}) {width_in:.2f}x{height_in:.2f} in]\n")
                for p in shape.text_frame.paragraphs:
                    text = p.text.strip()
                    if text:
                        font_info = []
                        for run in p.runs:
                            f_size = f"{run.font.size.pt:.1f}pt" if run.font and run.font.size else "default"
                            f_bold = "B" if run.font and run.font.bold else ""
                            font_info.append(f"{f_size}{f_bold}")
                        fi_str = f" [{', '.join(font_info)}]" if font_info else ""
                        out.write(f"    P: {text}{fi_str}\n")
            elif shape.has_table:
                table = shape.table
                out.write(f"\n  [Table {s_idx}: ({len(table.rows)}x{len(table.columns)}) at ({left_in:.2f}, {top_in:.2f}) {width_in:.2f}x{height_in:.2f} in]\n")
                for r_idx, row in enumerate(table.rows):
                    row_vals = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(f"    Row {r_idx}: {' | '.join(row_vals)}\n")
            else:
                out.write(f"\n  [Shape {s_idx}: {shape.name} (type: {shape.shape_type}) at ({left_in:.2f}, {top_in:.2f}) {width_in:.2f}x{height_in:.2f} in]\n")

print(f"Inspection written to {out_txt_path}")
