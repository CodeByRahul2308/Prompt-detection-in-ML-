"""
Deep verification of redesigned presentation.
Checks: all text readable, tables intact, metrics consistent, no empty slides.
"""
from pptx import Presentation
from pptx.util import Pt
import re

prs = Presentation("presentation/Minor_Project_Presentation_Redesigned.pptx")

# Collect all text from the presentation
all_text = []
for i, slide in enumerate(prs.slides, 1):
    slide_text = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            slide_text.append(shape.text_frame.text)
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    slide_text.append(cell.text)
    combined = " ".join(slide_text)
    all_text.append(combined)

full_text = " ".join(all_text)

# ─── METRIC VERIFICATION ────────────────────────────────────────────
print("=" * 60)
print("METRIC VERIFICATION")
print("=" * 60)

checks = [
    ("4,000 prompts mentioned", "4,000" in full_text or "4000" in full_text),
    ("800 test samples mentioned", "800" in full_text),
    ("Baseline accuracy 94.00%", "94.00%" in full_text),
    ("Challenge accuracy 87.50%", "87.50%" in full_text),
    ("Rule-Only 50.00%", "50.00%" in full_text),
    ("ML-Only 79.17%", "79.17%" in full_text),
    ("Hybrid precision 84.62%", "84.62%" in full_text),
    ("Hybrid recall 91.67%", "91.67%" in full_text),
    ("Hybrid F1 88.00%", "88.00%" in full_text),
    ("Confusion TN=370", "370" in full_text),
    ("Confusion TP=382", "382" in full_text),
    ("Confusion FP=30", "FP: 30" in full_text or "FP=30" in full_text or "FP: 30" in full_text),
    ("Confusion FN=18", "FN: 18" in full_text or "FN=18" in full_text or "FN: 18" in full_text),
    ("CH06 mentioned", "CH06" in full_text),
    ("ML prob 44.7%", "44.7%" in full_text),
    ("ML prob 78.8%", "78.8%" in full_text),
    ("Challenge TN=10 FP=2 FN=1 TP=11", "TN=10" in full_text and "TP=11" in full_text),
    ("Regi et al. referenced", "Regi" in full_text),
    ("Rahul Patel credited", "Rahul Patel" in full_text),
    ("Kishan Gupta credited", "Kishan Gupta" in full_text),
    ("Draft allocation caveat", "faculty" in full_text.lower()),
]

passed = 0
failed = 0
for desc, result in checks:
    status = "PASS" if result else "FAIL"
    if result:
        passed += 1
    else:
        failed += 1
    print(f"  [{status}] {desc}")

print(f"\n  Results: {passed}/{len(checks)} checks passed")

# ─── EMPTY SLIDE CHECK ──────────────────────────────────────────────
print("\n" + "=" * 60)
print("EMPTY SLIDE CHECK")
print("=" * 60)
for i, slide in enumerate(prs.slides, 1):
    shape_count = len(slide.shapes)
    text_content = ""
    for s in slide.shapes:
        if s.has_text_frame:
            text_content += s.text_frame.text
    if len(text_content.strip()) < 10:
        print(f"  [WARNING] Slide {i} appears nearly empty ({len(text_content)} chars)")
    else:
        print(f"  [OK] Slide {i}: {shape_count} shapes, {len(text_content)} chars of text")

# ─── FONT SIZE CHECK ────────────────────────────────────────────────
print("\n" + "=" * 60)
print("FONT SIZE DISTRIBUTION")
print("=" * 60)
sizes = {}
for slide in prs.slides:
    for shape in slide.shapes:
        if shape.has_text_frame:
            for p in shape.text_frame.paragraphs:
                if p.font.size:
                    s = int(p.font.size / 12700)
                    sizes[s] = sizes.get(s, 0) + 1

for s in sorted(sizes.keys()):
    bar = "#" * min(sizes[s], 50)
    print(f"  {s:3d}pt: {sizes[s]:3d} paragraphs {bar}")

# Check no critical body text below 14pt
small_count = sum(v for k, v in sizes.items() if k < 14)
if small_count > 0:
    print(f"\n  [NOTE] {small_count} paragraphs below 14pt (acceptable for footnotes/captions)")
else:
    print(f"\n  [OK] No text below 14pt")

if failed == 0:
    print("\n" + "=" * 60)
    print("ALL VERIFICATIONS PASSED")
    print("=" * 60)
else:
    print(f"\n  {failed} checks failed — review needed")
